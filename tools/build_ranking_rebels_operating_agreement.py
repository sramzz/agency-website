from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path("deliverables/Ranking_Rebels_LLC_Operating_Agreement_Bilingual.docx")


ENGLISH = [
    ("Article 1 Formation and Purpose", [
        ("1.1 Formation", "The Members adopt this Operating Agreement for Ranking Rebels LLC, a Wyoming limited liability company. The Company will come into existence when its Articles of Organization become effective under the Wyoming Limited Liability Company Act. If this Agreement is signed before that date, it becomes effective automatically on the effective date of the Articles of Organization, shown as [FORMATION DATE]."),
        ("1.2 Name", "The name of the Company is Ranking Rebels LLC. The Company may use assumed names, brands, and trade names approved by a Majority Vote and registered when required."),
        ("1.3 Purpose", "The Company may provide marketing, advertising, search engine optimization, consulting, software development, artificial intelligence, automation, and related services; acquire or invest in other businesses or assets; and conduct any other lawful activity permitted for a Wyoming limited liability company."),
        ("1.4 Duration", "The Company has perpetual duration unless dissolved under this Agreement or applicable law."),
        ("1.5 Offices and Registered Agent", "The principal office is [PRINCIPAL OFFICE ADDRESS]. The registered agent and registered office are [REGISTERED AGENT LEGAL NAME AND WYOMING ADDRESS]. A Majority Vote may change either address or the registered agent without amending this Agreement."),
        ("1.6 Governing Law", "Wyoming law governs the Company and this Agreement, without regard to conflict-of-laws rules, except where federal law or another mandatory law applies."),
    ]),
    ("Article 2 Members and Ownership", [
        ("2.1 Initial Members", "The initial Members, their contact information, Percentage Interests, and acknowledged contributions appear only in Schedule A. Each initial Member owns a 33 1/3 percent Percentage Interest immediately upon formation. The interests are fully issued and are not subject to vesting."),
        ("2.2 Nature of Contributions", "The descriptions in Schedule A acknowledge the experience, relationships, access, services, and property each Member expects to contribute. Except for the express assignment of the Domain Name under Exhibit B, those descriptions do not create a minimum-hours obligation, a guaranteed result, a separate capital credit, or additional compensation. No prospective customer is a Company asset until the customer signs a contract with the Company."),
        ("2.3 No Required Additional Contributions", "No Member must contribute additional cash, property, or services. Additional contributions are voluntary and do not change Percentage Interests unless all Members approve the change in a signed writing."),
        ("2.4 No Interest on Capital", "The Company will not pay interest on a contribution or capital account unless all Members approve a written loan on commercially reasonable terms."),
        ("2.5 Limited Liability", "No Member is personally liable for a Company debt or obligation solely because that person is a Member. A Member remains responsible for that Member's own fraud, willful misconduct, contractual guarantees, and obligations imposed by nonwaivable law."),
        ("2.6 No Guaranteed Employment", "Membership does not create employment, a minimum workload, a salary, or continued engagement. The Company may enter separate employment or service agreements approved under this Agreement."),
    ]),
    ("Article 3 Member Management and Voting", [
        ("3.1 Member Managed Company", "The Company is member-managed. Subject to this Agreement, every Member may participate in management and has voting power equal to that Member's Percentage Interest."),
        ("3.2 Majority Vote", "A Majority Vote means approval by Members holding more than 50 percent of all Percentage Interests entitled to vote. With the initial ownership, any two Members constitute a Majority Vote. Except for a Reserved Matter or a matter expressly governed by another standard, Company decisions require a Majority Vote."),
        ("3.3 Notice and Participation", "Every Member must receive reasonable notice and an opportunity to participate before a vote. Meetings may occur by video, telephone, or another real-time electronic method. Unless all Members waive notice, at least two business days' notice is required for an ordinary meeting and five business days for a meeting involving a Reserved Matter."),
        ("3.4 Written and Electronic Consent", "A decision may be approved without a meeting through a writing or electronic message that clearly identifies the decision and the approving Member. Consents may be signed in counterparts. The Company must retain the consent with its records."),
        ("3.5 Reserved Matters", "The following actions require the unanimous written approval of every then-current Member: amending the Articles of Organization or this Agreement; changing any Percentage Interest; issuing, granting, redeeming, or repurchasing an interest except under an existing express option in this Agreement; admitting a Member; waiving preemptive, first-refusal, or tag-along rights; borrowing money, guaranteeing an obligation, or granting a lien; acquiring another business or a material ownership interest in another entity; merging, converting, domesticating, or reorganizing the Company; selling all or substantially all Company assets; selling or exclusively licensing material Company intellectual property outside the ordinary course; approving or changing Member compensation; entering a material transaction with a Member or affiliate; changing federal tax classification; making a material tax election outside the ordinary course; dissolving the Company; or agreeing to take any of these actions."),
        ("3.6 Conflicts of Interest", "A Member must disclose any personal interest in a proposed Company transaction. A related-party transaction requires the approval stated in Section 3.5 and must be fair to the Company. Approval does not excuse fraud, concealment, or misuse of Company property."),
        ("3.7 Deadlock", "If the Members cannot obtain the approval required for a Reserved Matter, they will meet in good faith and then use the mediation procedure in Section 13.2. If mediation does not resolve the matter, the status quo continues. Deadlock alone does not force a sale, buyout, or dissolution."),
    ]),
    ("Article 4 Authority Spending and Emergencies", [
        ("4.1 Individual Spending Authority", "A Member may approve and incur an ordinary-course Company expense without a vote only when it is within an approved budget and does not exceed the lesser of US$1,500 or 15 percent of Current Monthly Revenue. Current Monthly Revenue means gross revenue actually collected during the immediately preceding full calendar month, excluding taxes collected for a government, refunds, and chargebacks. Related purchases must be aggregated and may not be divided to avoid this limit."),
        ("4.2 Startup Spending", "Before the Company has Current Monthly Revenue, or when 15 percent of Current Monthly Revenue is zero, every Company expense requires a Majority Vote. The Members may approve a written startup budget by Majority Vote; an individual Member may spend within a specific line item of that budget up to the amount approved."),
        ("4.3 Contracts", "Any Member may sign a client or vendor contract previously approved by a Majority Vote. A Member may sign an unapproved ordinary client contract only if it uses Company-approved terms and does not create an expense, guarantee, exclusivity obligation, transfer of Company intellectual property, uncapped liability, or term longer than twelve months. Any other contract requires a Majority Vote or unanimous approval if it concerns a Reserved Matter."),
        ("4.4 No Unauthorized Commitments", "A Member may not bind the Company outside the authority granted by this Agreement or a documented approval. The Member must promptly disclose any commitment made for the Company and provide the contract, invoice, or other supporting record."),
        ("4.5 Emergency Action", "An Emergency exists only when action is reasonably required within 48 hours to prevent material financial, legal, cybersecurity, operational, or client harm. Two Members may take the minimum reasonable Emergency action after making a documented good-faith attempt to contact the third Member by both email and the designated group messaging channel. They must report the action and supporting facts promptly. The Emergency procedure never authorizes a Reserved Matter involving ownership, admission or dilution, borrowing, acquisition, sale of the Company, tax classification, amendment of this Agreement, or dissolution."),
        ("4.6 Banking Controls", "Company funds must remain in accounts titled in the Company's name. No Member may mix Company funds with personal funds. Banking access and payment controls require a Majority Vote, but no control may expand an individual's spending authority under this Article."),
    ]),
    ("Article 5 Compensation Accounting and Distributions", [
        ("5.1 Member Compensation", "A Member may receive salary, contractor fees, commissions, or other compensation only under a written arrangement approved unanimously. Compensation is a Company expense before Net Profits are determined and does not change Percentage Interests unless all Members expressly agree in writing."),
        ("5.2 Reimbursement", "The Company will reimburse a Member for a documented, reasonable Company expense authorized under this Agreement. Reimbursement is not compensation or a capital contribution."),
        ("5.3 Books and Accounting", "The Company will maintain complete and accurate books in United States dollars using a method selected by Majority Vote after consultation with its tax adviser. The fiscal and tax year ends December 31 unless law or a unanimous tax decision requires otherwise."),
        ("5.4 Allocations", "Subject to applicable tax law, Net Profits and Net Losses will be allocated among the Members in proportion to their Percentage Interests. The Company's tax adviser may make technical allocations required by Internal Revenue Code section 704 and its regulations, provided the adjustment does not intentionally alter the agreed economic arrangement."),
        ("5.5 Distributions", "After paying or reserving for expenses, debt, required tax withholding, and reasonable working capital, the Members may authorize cash distributions by Majority Vote. Distributions will be made in proportion to Percentage Interests, subject to legally required withholding and any specific offset permitted by this Agreement."),
        ("5.6 Withholding", "The Company may withhold and remit any tax required with respect to a Member, including withholding concerning a foreign Member. An amount paid to a tax authority for a Member is treated as distributed to that Member to the extent permitted by tax law. Each Member must timely provide Forms W-8, taxpayer-identification information to the tax adviser, and other certifications required for compliance. Taxpayer identification numbers will not be inserted into this Agreement."),
        ("5.7 Restrictions on Distributions", "No distribution or redemption may be made when prohibited by the Wyoming Limited Liability Company Act or other applicable law. A prohibited payment will be postponed, not forgiven, until it can lawfully be made."),
    ]),
    ("Article 6 Tax Matters", [
        ("6.1 Default Classification", "Unless all Members approve a different classification and the required election is filed, the Company will be treated as a partnership for United States federal income tax purposes."),
        ("6.2 Returns and Information", "The Company will obtain an employer identification number and engage a qualified United States tax professional. It will prepare and timely file required returns and provide Members with available tax information reasonably soon after each year-end. Each Member is responsible for that Member's personal filings and taxes in every applicable jurisdiction."),
        ("6.3 Foreign Member Compliance", "The Members acknowledge that foreign ownership, services performed in multiple countries, United States clients, and transfers of Company interests may create reporting and withholding obligations. The Company may retain funds and take actions reasonably recommended by its tax adviser to comply with those obligations."),
        ("6.4 Partnership Representative", "For each tax year in which a partnership representative is required, the Members will unanimously designate an eligible person with the substantial United States presence required by federal law. Until designated on the applicable return, the position is [UNITED STATES PARTNERSHIP REPRESENTATIVE]. The representative must keep the Members informed and obtain unanimous approval before settling an audit, extending a limitation period, making a push-out election, or taking another material position, except when immediate action is legally required. These internal limits do not restrict the representative's authority against the Internal Revenue Service."),
        ("6.5 Tax Elections", "A tax classification election and any election that materially changes the Members' economic or filing position require unanimous approval. Routine elections and return positions consistent with this Agreement may be made by Majority Vote upon advice of the Company's tax professional."),
        ("6.6 No Tax Assurance", "No Member or the Company guarantees a particular tax result. Each Member must obtain independent advice concerning United States, Australian, Dutch, Colombian, and other applicable law."),
    ]),
    ("Article 7 Company Property Technology and Confidentiality", [
        ("7.1 Company Property", "All money, contracts, receivables, customer information, work product, accounts, records, and other assets acquired for the Company belong solely to the Company. No Member has a direct ownership claim to a specific Company asset."),
        ("7.2 Domain and Brand", "SR will assign rankingrebels.com and the associated goodwill to the Company under Exhibit B immediately after formation and will complete the registrar transfer when technically available. All registered and unregistered Ranking Rebels names, marks, designs, content, and goodwill created for the business belong to the Company."),
        ("7.3 Work Product", "To the fullest extent permitted by law, each Member assigns to the Company all worldwide rights in work product created specifically for the Company or its customers, including software, prompts, automations, designs, copy, campaigns, processes, documentation, data structures, and inventions. General skills, experience, and independently developed material unrelated to Company business remain with the Member. A Member must identify any pre-existing material before incorporating it and grant the Company a perpetual, worldwide, transferable, sublicensable, royalty-free license necessary to use the resulting work product."),
        ("7.4 Further Assurances", "Each Member will sign documents and take reasonable steps necessary to confirm, register, enforce, or transfer Company intellectual property. To the extent legally permitted, each Member waives and agrees not to assert moral rights in Company work product."),
        ("7.5 Hetzner Hosting", "SR may initially host Company systems on SR's Hetzner virtual private server. The Company will reimburse documented third-party hosting costs actually attributable to Company use, without markup, under the approval and reimbursement rules of this Agreement. Hosting does not change ownership. The Company must receive current administrative access, credentials, export capability, and reasonably current backups. SR must cooperate with a migration selected by Majority Vote and may not suspend access to pressure the Company or another Member. Following a request to migrate or the end of SR's involvement, SR will provide reasonable transition assistance for up to 30 days."),
        ("7.6 Confidential Information", "Confidential Information includes nonpublic business plans, pricing, finances, credentials, source code, automations, prompts, customer and prospect information, strategies, analytics, trade secrets, personal data, and this Agreement's nonpublic schedules. A Member may use Confidential Information only for Company purposes and may disclose it only to authorized persons under suitable confidentiality obligations or as legally required after giving prompt notice when permitted."),
        ("7.7 Security and Return", "Each Member must use reasonable security measures, report a suspected breach promptly, and return or securely delete Company information and credentials when requested or upon dissociation, except for a legally required archival copy."),
        ("7.8 Survival", "Confidentiality continues indefinitely for trade secrets and for five years after dissociation for other Confidential Information, or longer where mandatory law or a customer contract requires."),
    ]),
    ("Article 8 Issuance Admission and Transfers", [
        ("8.1 New Interests and Members", "The Company may issue an interest or admit a new Member only with unanimous written approval. A new Member must sign the joinder in Exhibit D or a substantially similar document. A transferee receives only economic rights unless unanimously admitted as a Member."),
        ("8.2 Preemptive Rights", "Before issuing a new interest, option, warrant, or convertible right, the Company must offer each Member the right to purchase that Member's pro-rata portion on the same material terms. The offer remains open for 20 business days. Issuances under a unanimously approved transaction may state different procedures, but waiver of a Member's right requires that Member's written consent."),
        ("8.3 Transfer Restriction", "A Member may not sell, assign, gift, pledge, encumber, or otherwise transfer any part of an interest except under this Agreement. An attempted unauthorized transfer is ineffective to the fullest extent permitted by law."),
        ("8.4 Right of First Refusal", "Before a voluntary transfer to an outsider, the transferring Member must deliver the complete bona fide written offer and proposed terms to the Company and other Members. The Company has 20 business days to purchase all or part on those terms. The remaining Members then have 15 business days to purchase the remainder pro rata or as they agree. If the rights are not fully exercised, the transfer may close within 60 days on terms no more favorable to the outsider, subject to the tag-along right and the rule that the buyer is not a Member without unanimous admission."),
        ("8.5 Tag Along Right", "If an outsider proposes to purchase any interest from a Member, each other Member may include the same proportion of that other Member's interest in the sale on identical per-percentage economic terms. The selling Member must ensure the buyer purchases the included interests or may not complete the proposed transfer."),
        ("8.6 No Forced Sale by Majority", "No drag-along right applies. A sale of all interests or all or substantially all Company assets requires unanimous approval."),
        ("8.7 Compliance and Withholding", "Every transfer is subject to applicable securities, tax, sanctions, anti-money-laundering, and other laws. The Company and a buyer may withhold amounts required by law, including withholding connected with a foreign person's transfer of a partnership interest."),
    ]),
    ("Article 9 Dissociation Removal and Buyouts", [
        ("9.1 Voluntary Departure", "A Member may resign from active management by giving at least 90 days' written notice. On the effective date, the Member becomes a dissociated economic-interest holder and no longer votes or acts for the Company, but retains economic rights until the interest is purchased or otherwise transferred."),
        ("9.2 Company and Member Purchase Options", "Following voluntary departure, death, Permanent Incapacity, or removal for Cause, the Company has the first option to purchase all of the affected interest. The Company may exercise within 60 days after receiving notice of the event or the final valuation, whichever is later. If the Company does not purchase all of it, the remaining Members have a further 30 days to purchase the balance pro rata or as they agree."),
        ("9.3 Death and Heirs", "Upon a Member's death, the personal representative or heirs receive only the economic rights associated with the interest and do not become Members. They must sell the interest if the Company or remaining Members timely exercise the purchase option. Until closing, the holder receives distributions attributable to that interest, subject to withholding and offsets."),
        ("9.4 Permanent Incapacity", "Permanent Incapacity means an inability, because of physical or mental condition, to participate materially in Company management for 180 consecutive days, supported by reasonable medical evidence that respects applicable privacy law. Permanent Incapacity creates a purchase option but does not constitute misconduct."),
        ("9.5 Cause", "Cause means fraud, theft, embezzlement, intentional misconduct materially harmful to the Company, knowing misuse of Company funds or credentials, a material breach of confidentiality or intellectual-property duties, conviction of a serious dishonesty offense, or another material breach of this Agreement that remains uncured after 15 business days' written notice when cure is possible."),
        ("9.6 Removal Procedure", "The two unaffected initial Members may remove a Member from management for Cause. The affected Member must receive detailed written notice and at least five business days to respond before the decision, unless immediate suspension of access is reasonably necessary to protect funds, systems, evidence, or customers. The affected Member may present information but may not vote. Removal ends management and voting rights but does not confiscate economic ownership. The Company may require a sale under Section 9.2."),
        ("9.7 Fair Market Value", "The purchase price equals the affected Percentage Interest's proportionate share of the Company's fair market value as a going concern, without minority or lack-of-marketability discounts, less only an offset agreed in writing or established by a final judgment or arbitration award. The Company and seller will attempt to agree on value within 20 business days. If they do not, they will jointly appoint one independent qualified business appraiser. If they cannot agree on an appraiser within ten business days, the ICDR will appoint one upon either party's request. The appraiser's determination is binding absent manifest error, and the Company and seller share the appraisal cost equally."),
        ("9.8 Payment Terms", "Unless the parties agree otherwise, the buyer pays 20 percent at closing and the balance in 24 equal monthly installments under a promissory note. Interest accrues at the United States prime rate published by The Wall Street Journal on the closing date plus two percentage points, not exceeding the lawful maximum. The buyer may prepay without penalty. No Member gives a personal guarantee solely because of the purchase. Any payment prohibited by law is postponed until lawful and continues to accrue interest."),
        ("9.9 Closing", "Closing will occur within 30 days after price determination, subject to legally required withholding. At closing, the seller transfers the interest free of liens, resigns from Company positions, returns Company property, and signs reasonable transfer documents. The buyer delivers the down payment and promissory note."),
        ("9.10 Continuation", "A Member's dissociation, death, incapacity, bankruptcy, or transfer does not dissolve the Company."),
    ]),
    ("Article 10 Duties and Protection", [
        ("10.1 Standards of Conduct", "Each Member must act in good faith, deal fairly with the Company and other Members, avoid intentional harm, disclose material conflicts, safeguard Company property, and comply with the duties that cannot lawfully be waived under Wyoming law."),
        ("10.2 Reliance", "A Member acting in good faith may rely on Company records and on information, opinions, and reports from qualified professionals reasonably believed to be reliable."),
        ("10.3 Indemnification", "To the fullest extent permitted by law, the Company will indemnify a Member for a claim or expense arising from authorized Company activity performed in good faith. No indemnification applies to fraud, theft, willful misconduct, a knowing violation of law, an unauthorized personal benefit, or a material breach of this Agreement."),
        ("10.4 Advancement", "The Company may advance defense expenses after receiving a written promise to repay them if indemnification is ultimately unavailable. Advancement requires approval of all disinterested Members."),
        ("10.5 Insurance", "The Company may purchase liability, cyber, errors-and-omissions, key-person, or other insurance by Majority Vote."),
    ]),
    ("Article 11 Nonsolicitation and Business Protections", [
        ("11.1 No Broad Noncompetition Covenant", "This Agreement does not broadly prohibit a former Member from earning a living or operating another business. The specific confidentiality, intellectual-property, and nonsolicitation duties below protect legitimate Company interests."),
        ("11.2 Client Nonsolicitation", "During membership and for 24 months after dissociation, a former Member will not knowingly solicit or induce a Company client to end or materially reduce its relationship with the Company when the former Member had material contact with that client or received material Confidential Information about it during the preceding twelve months."),
        ("11.3 Workforce Nonsolicitation", "During membership and for 24 months after dissociation, a former Member will not knowingly solicit a Company employee or contractor with whom the former Member worked materially during the preceding twelve months to terminate that relationship. General advertising not targeted at such a person and hiring someone who responds without targeted solicitation do not violate this Section."),
        ("11.4 Lawful Scope", "Sections 11.2 and 11.3 apply only to the maximum extent enforceable under the law applicable to the affected person and conduct. A court or arbitrator may narrow an invalid duration, activity, or scope rather than invalidate more of the protection than necessary. Nothing restricts legally protected reporting, competition, or use of general skills."),
        ("11.5 Injunctive Relief", "Actual or threatened misuse of Company intellectual property, trade secrets, credentials, or Confidential Information may cause harm not adequately repaired by money. The Company may seek temporary or injunctive relief as allowed by Section 13.4 without waiving arbitration."),
    ]),
    ("Article 12 Records and Compliance", [
        ("12.1 Records", "The Company will maintain formation documents, this Agreement and amendments, Member consents, financial statements, tax returns, material contracts, ownership records, and other records required by law. Electronic records are permitted."),
        ("12.2 Inspection", "A Member may inspect reasonable Company records for a purpose related to membership upon reasonable notice, subject to confidentiality, privilege, privacy, and security safeguards. A dissociated holder receives only information required by law or reasonably necessary to verify distributions and a buyout."),
        ("12.3 Separate Accounts and Documentation", "The Members will preserve the Company's separate legal existence, use Company accounts and contracts, document material approvals, and avoid representing that Company obligations are personal obligations unless a separate guarantee is intentionally signed."),
        ("12.4 Multijurisdictional Compliance", "Before opening an office, hiring personnel, creating a regular place of business, or taking another step likely to require registration or tax presence in a country or state, the Members will obtain appropriate advice and approve the step under the voting standard otherwise applicable. Each Member will accurately report where services are performed."),
    ]),
    ("Article 13 Dispute Resolution", [
        ("13.1 Good Faith Discussion", "A party must first give written notice describing a dispute and requested remedy. The parties will meet remotely within ten business days and attempt in good faith to resolve it."),
        ("13.2 Mediation", "If unresolved, the parties will participate in confidential remote mediation with a mutually selected mediator. If they cannot select one within ten business days, the ICDR will provide or appoint a mediator. The parties share mediator fees equally unless they agree otherwise. Mediation ends 30 days after appointment unless extended in writing."),
        ("13.3 Arbitration", "A dispute not resolved by mediation will be finally resolved by one arbitrator under the International Dispute Resolution Procedures of the International Centre for Dispute Resolution, including its applicable commercial rules. The legal seat is Cheyenne, Wyoming. Hearings will be remote unless the arbitrator determines an in-person hearing is necessary. The arbitration language is English. The award may grant legal or equitable relief and may allocate fees and reasonable legal costs. Judgment may be entered in any court with jurisdiction."),
        ("13.4 Emergency Court Relief", "A party may seek temporary relief from a court with jurisdiction to protect funds, systems, credentials, evidence, Confidential Information, intellectual property, or customer data pending arbitration. Doing so does not waive mediation or arbitration."),
        ("13.5 Confidentiality", "The parties will keep mediation and arbitration confidential except as necessary to enforce an award, obtain professional advice, comply with law, or protect a legal right."),
    ]),
    ("Article 14 Dissolution and Winding Up", [
        ("14.1 Dissolution Events", "The Company dissolves upon unanimous written approval, an event that makes the business unlawful, or a final judicial decree requiring dissolution. Deadlock by itself is not a dissolution event."),
        ("14.2 Winding Up", "A person unanimously selected by the Members will wind up the Company. The Company will collect assets, complete or terminate contracts, pay or reserve for liabilities, and distribute any remainder as required by law and this Agreement."),
        ("14.3 Final Distribution", "After creditors and required reserves are satisfied, remaining assets will be distributed to the Members and economic-interest holders in accordance with positive capital accounts when tax law requires and otherwise in proportion to their Percentage Interests, subject to lawful offsets and withholding."),
    ]),
    ("Article 15 General Terms", [
        ("15.1 Notices", "Formal notices must be sent to the mailing and email addresses in Schedule A, as updated by written notice. Email is effective when sent without a delivery-failure notice. A notice concerning an Emergency must also be sent through the Company's designated group messaging channel. Legal service of process must follow applicable law."),
        ("15.2 Amendments", "This Agreement may be amended only by a writing signed by every Member. No course of dealing or electronic discussion amends the Agreement unless it clearly states an intent to amend and receives every required signature."),
        ("15.3 Entire Agreement", "This Agreement, its schedules and exhibits, and any later signed joinder or amendment constitute the entire agreement among the Members concerning Company governance and ownership. They supersede inconsistent prior oral or written understandings."),
        ("15.4 Priority", "As between the Members, this Agreement controls to the fullest extent permitted by law. A filed Company record controls only where applicable law gives it priority. A mandatory rule of law controls over an inconsistent provision."),
        ("15.5 Severability", "If a provision is unenforceable, it will be enforced to the maximum lawful extent or severed, and the remainder will continue. A court or arbitrator may reform a restriction only as permitted by applicable law."),
        ("15.6 Waiver", "A waiver must be written and applies only to the stated instance. Delay or partial exercise of a right is not a waiver."),
        ("15.7 No Third Party Beneficiaries", "Except for an indemnified person and a permitted successor, this Agreement creates no right in a nonparty."),
        ("15.8 Counterparts and Electronic Signatures", "This Agreement may be signed in counterparts and by a reliable electronic-signature method. All counterparts form one instrument."),
        ("15.9 Further Documents", "Each Member will sign reasonable documents needed to carry out this Agreement, including tax, banking, transfer, and intellectual-property documents."),
        ("15.10 English Controls", "The Spanish version is provided for convenience and mutual understanding. If the versions differ, the English version controls."),
    ]),
]


SPANISH = [
    ("Artículo 1 Constitución y Objeto", [
        ("1.1 Constitución", "Los Miembros adoptan este Acuerdo Operativo para Ranking Rebels LLC, una compañía de responsabilidad limitada de Wyoming. La Compañía existirá cuando sus Artículos de Organización entren en vigor conforme a la Ley de Compañías de Responsabilidad Limitada de Wyoming. Si este Acuerdo se firma antes de esa fecha, entrará en vigor automáticamente en la fecha efectiva de los Artículos de Organización, indicada como [FECHA DE CONSTITUCIÓN]."),
        ("1.2 Nombre", "El nombre de la Compañía es Ranking Rebels LLC. La Compañía podrá utilizar nombres comerciales, marcas y nombres ficticios aprobados por Voto Mayoritario y registrados cuando sea obligatorio."),
        ("1.3 Objeto", "La Compañía podrá prestar servicios de marketing, publicidad, optimización para motores de búsqueda, consultoría, desarrollo de software, inteligencia artificial, automatización y servicios relacionados; adquirir o invertir en otros negocios o activos; y realizar cualquier otra actividad lícita permitida a una compañía de responsabilidad limitada de Wyoming."),
        ("1.4 Duración", "La Compañía tendrá duración perpetua, salvo que se disuelva conforme a este Acuerdo o la ley aplicable."),
        ("1.5 Oficinas y Agente Registrado", "La oficina principal será [DIRECCIÓN DE LA OFICINA PRINCIPAL]. El agente registrado y la oficina registrada serán [NOMBRE LEGAL DEL AGENTE REGISTRADO Y DIRECCIÓN EN WYOMING]. Un Voto Mayoritario podrá cambiar cualquiera de estas direcciones o el agente registrado sin modificar este Acuerdo."),
        ("1.6 Ley Aplicable", "Las leyes de Wyoming regirán la Compañía y este Acuerdo, sin aplicar sus normas sobre conflicto de leyes, salvo cuando se aplique obligatoriamente una ley federal u otra norma imperativa."),
    ]),
    ("Artículo 2 Miembros y Propiedad", [
        ("2.1 Miembros Iniciales", "Los Miembros iniciales, sus datos de contacto, Participaciones Porcentuales y aportes reconocidos figuran únicamente en el Anexo A. Cada Miembro inicial será propietario de una Participación Porcentual de 33 1/3 por ciento desde la constitución. Las participaciones se emiten en su totalidad y no están sujetas a consolidación gradual."),
        ("2.2 Naturaleza de los Aportes", "Las descripciones del Anexo A reconocen la experiencia, relaciones, acceso, servicios y bienes que cada Miembro espera aportar. Salvo la cesión expresa del Nombre de Dominio conforme al Anexo B, dichas descripciones no crean una obligación de horas mínimas, un resultado garantizado, un crédito de capital separado ni una compensación adicional. Ningún cliente potencial será activo de la Compañía hasta que firme un contrato con ella."),
        ("2.3 Sin Aportes Adicionales Obligatorios", "Ningún Miembro estará obligado a aportar dinero, bienes o servicios adicionales. Los aportes adicionales serán voluntarios y no cambiarán las Participaciones Porcentuales salvo que todos los Miembros aprueben el cambio mediante escrito firmado."),
        ("2.4 Sin Intereses sobre el Capital", "La Compañía no pagará intereses sobre aportes o cuentas de capital salvo que todos los Miembros aprueben por escrito un préstamo en condiciones comercialmente razonables."),
        ("2.5 Responsabilidad Limitada", "Ningún Miembro será personalmente responsable de una deuda u obligación de la Compañía por el solo hecho de ser Miembro. Cada Miembro seguirá siendo responsable de su propio fraude, conducta dolosa, garantías contractuales y obligaciones impuestas por normas que no puedan renunciarse."),
        ("2.6 Sin Empleo Garantizado", "La condición de Miembro no crea una relación laboral, carga mínima de trabajo, salario ni contratación continuada. La Compañía podrá celebrar contratos laborales o de servicios separados aprobados conforme a este Acuerdo."),
    ]),
    ("Artículo 3 Administración y Votación de los Miembros", [
        ("3.1 Compañía Administrada por sus Miembros", "La Compañía será administrada por sus Miembros. Sujeto a este Acuerdo, cada Miembro podrá participar en la administración y tendrá poder de voto equivalente a su Participación Porcentual."),
        ("3.2 Voto Mayoritario", "Voto Mayoritario significa la aprobación de Miembros que posean más del 50 por ciento de todas las Participaciones Porcentuales con derecho a voto. Con la propiedad inicial, cualesquiera dos Miembros constituyen un Voto Mayoritario. Salvo para una Materia Reservada o una materia sujeta expresamente a otra regla, las decisiones de la Compañía requerirán Voto Mayoritario."),
        ("3.3 Aviso y Participación", "Cada Miembro deberá recibir aviso razonable y oportunidad de participar antes de una votación. Las reuniones podrán celebrarse por video, teléfono u otro medio electrónico en tiempo real. Salvo renuncia de todos los Miembros al aviso, se requerirán al menos dos días hábiles para una reunión ordinaria y cinco días hábiles para una reunión sobre una Materia Reservada."),
        ("3.4 Consentimiento Escrito y Electrónico", "Una decisión podrá aprobarse sin reunión mediante escrito o mensaje electrónico que identifique claramente la decisión y al Miembro que la aprueba. Los consentimientos podrán firmarse en ejemplares separados. La Compañía conservará el consentimiento en sus registros."),
        ("3.5 Materias Reservadas", "Las siguientes acciones requerirán aprobación escrita unánime de todos los Miembros existentes: modificar los Artículos de Organización o este Acuerdo; cambiar cualquier Participación Porcentual; emitir, otorgar, rescatar o recomprar una participación, salvo bajo una opción expresa ya contenida en este Acuerdo; admitir un Miembro; renunciar a derechos de suscripción preferente, preferencia de compra o acompañamiento; tomar préstamos, garantizar obligaciones u otorgar gravámenes; adquirir otro negocio o una participación material en otra entidad; fusionar, convertir, trasladar o reorganizar la Compañía; vender todos o sustancialmente todos los activos; vender u otorgar licencia exclusiva sobre propiedad intelectual material fuera del curso ordinario; aprobar o cambiar la compensación de un Miembro; celebrar una transacción material con un Miembro o afiliado; cambiar la clasificación fiscal federal; realizar una elección fiscal material fuera del curso ordinario; disolver la Compañía; o comprometerse a realizar cualquiera de estas acciones."),
        ("3.6 Conflictos de Interés", "Cada Miembro deberá revelar cualquier interés personal en una transacción propuesta. Una transacción con parte relacionada requerirá la aprobación indicada en la Sección 3.5 y deberá ser justa para la Compañía. La aprobación no excusará fraude, ocultamiento ni uso indebido de bienes de la Compañía."),
        ("3.7 Bloqueo Decisorio", "Si no se obtiene la aprobación requerida para una Materia Reservada, los Miembros se reunirán de buena fe y luego seguirán la mediación de la Sección 13.2. Si la mediación no resuelve la materia, se mantendrá la situación existente. El bloqueo, por sí solo, no obligará a una venta, compra de participación ni disolución."),
    ]),
    ("Artículo 4 Autoridad Gastos y Emergencias", [
        ("4.1 Autoridad Individual de Gasto", "Un Miembro podrá aprobar e incurrir en un gasto ordinario sin votación únicamente si está dentro de un presupuesto aprobado y no supera el menor de US$1.500 o el 15 por ciento de los Ingresos Mensuales Actuales. Ingresos Mensuales Actuales significa los ingresos brutos efectivamente cobrados durante el mes calendario completo inmediatamente anterior, excluidos impuestos recaudados para una autoridad, reembolsos y contracargos. Las compras relacionadas se sumarán y no podrán dividirse para evitar este límite."),
        ("4.2 Gastos Iniciales", "Antes de que existan Ingresos Mensuales Actuales, o cuando su 15 por ciento sea cero, todo gasto requerirá Voto Mayoritario. Los Miembros podrán aprobar por Voto Mayoritario un presupuesto inicial escrito; un Miembro individual podrá gastar dentro de una partida específica hasta el monto aprobado."),
        ("4.3 Contratos", "Cualquier Miembro podrá firmar un contrato con cliente o proveedor previamente aprobado por Voto Mayoritario. Un Miembro podrá firmar un contrato ordinario con cliente no aprobado si utiliza condiciones aprobadas por la Compañía y no crea un gasto, garantía, exclusividad, transferencia de propiedad intelectual, responsabilidad ilimitada ni plazo superior a doce meses. Cualquier otro contrato requerirá Voto Mayoritario o aprobación unánime si corresponde a una Materia Reservada."),
        ("4.4 Sin Compromisos No Autorizados", "Ningún Miembro podrá obligar a la Compañía fuera de la autoridad conferida por este Acuerdo o una aprobación documentada. El Miembro revelará prontamente cualquier compromiso contraído para la Compañía y entregará el contrato, factura u otro respaldo."),
        ("4.5 Acción de Emergencia", "Existirá una Emergencia únicamente cuando sea razonablemente necesario actuar dentro de 48 horas para evitar un daño material financiero, legal, de ciberseguridad, operativo o a un cliente. Dos Miembros podrán tomar la medida mínima razonable después de intentar de buena fe y de forma documentada contactar al tercer Miembro por correo electrónico y por el canal grupal de mensajería designado. Informarán prontamente la medida y sus fundamentos. Este procedimiento nunca autorizará una Materia Reservada relativa a propiedad, admisión o dilución, endeudamiento, adquisición, venta de la Compañía, clasificación fiscal, modificación del Acuerdo o disolución."),
        ("4.6 Controles Bancarios", "Los fondos permanecerán en cuentas a nombre de la Compañía. Ningún Miembro mezclará fondos de la Compañía con fondos personales. El acceso bancario y los controles de pago requerirán Voto Mayoritario, sin ampliar la autoridad individual de gasto de este Artículo."),
    ]),
    ("Artículo 5 Compensación Contabilidad y Distribuciones", [
        ("5.1 Compensación de Miembros", "Un Miembro podrá recibir salario, honorarios, comisiones u otra compensación únicamente bajo un acuerdo escrito aprobado por unanimidad. La compensación será un gasto anterior al cálculo de Utilidades Netas y no cambiará las Participaciones Porcentuales salvo acuerdo escrito de todos los Miembros."),
        ("5.2 Reembolsos", "La Compañía reembolsará a un Miembro los gastos documentados y razonables autorizados conforme a este Acuerdo. El reembolso no será compensación ni aporte de capital."),
        ("5.3 Libros y Contabilidad", "La Compañía mantendrá libros completos y exactos en dólares estadounidenses mediante un método seleccionado por Voto Mayoritario después de consultar al asesor fiscal. El ejercicio fiscal y tributario terminará el 31 de diciembre, salvo que la ley o una decisión fiscal unánime exijan otra fecha."),
        ("5.4 Asignaciones", "Sujeto a la ley fiscal, las Utilidades Netas y Pérdidas Netas se asignarán según las Participaciones Porcentuales. El asesor fiscal podrá realizar asignaciones técnicas exigidas por la sección 704 del Código de Rentas Internas y sus reglamentos, siempre que no alteren intencionalmente el acuerdo económico."),
        ("5.5 Distribuciones", "Después de pagar o reservar para gastos, deudas, retenciones fiscales obligatorias y capital de trabajo razonable, los Miembros podrán autorizar distribuciones por Voto Mayoritario. Se realizarán según las Participaciones Porcentuales, sujetas a retenciones y compensaciones permitidas."),
        ("5.6 Retenciones", "La Compañía podrá retener y remitir todo impuesto obligatorio respecto de un Miembro, incluso por ser extranjero. Un monto pagado a una autoridad por cuenta de un Miembro se considerará distribuido a ese Miembro en la medida permitida por la ley fiscal. Cada Miembro entregará oportunamente formularios W-8, información de identificación tributaria al asesor fiscal y demás certificaciones necesarias. Los números de identificación tributaria no se incluirán en este Acuerdo."),
        ("5.7 Restricciones a las Distribuciones", "No se realizará distribución o rescate prohibido por la Ley de Compañías de Responsabilidad Limitada de Wyoming u otra ley aplicable. Un pago prohibido se aplazará, sin extinguirse, hasta que sea lícito."),
    ]),
    ("Artículo 6 Asuntos Fiscales", [
        ("6.1 Clasificación Predeterminada", "Salvo que todos los Miembros aprueben otra clasificación y se presente la elección correspondiente, la Compañía será tratada como sociedad de personas para fines del impuesto federal sobre la renta de Estados Unidos."),
        ("6.2 Declaraciones e Información", "La Compañía obtendrá un número de identificación del empleador y contratará a un profesional fiscal estadounidense calificado. Preparará y presentará oportunamente las declaraciones requeridas y entregará a los Miembros la información fiscal disponible razonablemente pronto después de cada cierre anual. Cada Miembro será responsable de sus declaraciones e impuestos personales en toda jurisdicción aplicable."),
        ("6.3 Cumplimiento de Miembros Extranjeros", "Los Miembros reconocen que la propiedad extranjera, los servicios en varios países, los clientes estadounidenses y las transferencias de participaciones pueden generar obligaciones de información y retención. La Compañía podrá retener fondos y tomar medidas razonablemente recomendadas por su asesor fiscal."),
        ("6.4 Representante Fiscal de la Sociedad", "Para cada año en que se exija un representante fiscal de la sociedad, los Miembros designarán por unanimidad a una persona elegible con la presencia sustancial en Estados Unidos exigida por la ley federal. Hasta su designación en la declaración aplicable, el cargo será [REPRESENTANTE FISCAL EN ESTADOS UNIDOS]. El representante informará a los Miembros y obtendrá aprobación unánime antes de resolver una auditoría, ampliar un plazo de prescripción, hacer una elección de traslado de ajustes o adoptar otra posición material, salvo que la ley exija acción inmediata. Estas limitaciones internas no restringen su autoridad frente al Servicio de Impuestos Internos."),
        ("6.5 Elecciones Fiscales", "Una elección de clasificación fiscal y cualquier elección que cambie materialmente la situación económica o declarativa de los Miembros requerirá aprobación unánime. Las elecciones rutinarias y posiciones coherentes con este Acuerdo podrán adoptarse por Voto Mayoritario con asesoría fiscal."),
        ("6.6 Sin Garantía Fiscal", "Ningún Miembro ni la Compañía garantiza un resultado fiscal. Cada Miembro deberá obtener asesoría independiente sobre las leyes de Estados Unidos, Australia, Países Bajos, Colombia y demás leyes aplicables."),
    ]),
    ("Artículo 7 Bienes Tecnología y Confidencialidad", [
        ("7.1 Bienes de la Compañía", "Todo dinero, contrato, cuenta por cobrar, información de clientes, producto de trabajo, cuenta, registro y demás activo adquirido para la Compañía pertenecerá únicamente a ella. Ningún Miembro tendrá propiedad directa sobre un activo específico."),
        ("7.2 Dominio y Marca", "SR cederá rankingrebels.com y su goodwill asociado a la Compañía conforme al Anexo B inmediatamente después de la constitución y completará la transferencia en el registrador cuando sea técnicamente posible. Todos los nombres, marcas, diseños, contenidos y goodwill registrados o no registrados creados para Ranking Rebels pertenecerán a la Compañía."),
        ("7.3 Producto de Trabajo", "En la máxima medida legal, cada Miembro cede a la Compañía todos los derechos mundiales sobre el producto de trabajo creado específicamente para la Compañía o sus clientes, incluidos software, prompts, automatizaciones, diseños, textos, campañas, procesos, documentación, estructuras de datos e invenciones. Las habilidades generales, experiencia y materiales desarrollados independientemente y no relacionados seguirán siendo del Miembro. El Miembro identificará materiales preexistentes antes de incorporarlos y otorgará a la Compañía una licencia perpetua, mundial, transferible, sublicenciable y libre de regalías necesaria para usar el resultado."),
        ("7.4 Documentos Adicionales", "Cada Miembro firmará documentos y realizará pasos razonables para confirmar, registrar, hacer valer o transferir propiedad intelectual. En la medida legal, renuncia a ejercer derechos morales sobre el producto de trabajo de la Compañía."),
        ("7.5 Alojamiento en Hetzner", "SR podrá alojar inicialmente los sistemas en su servidor privado virtual de Hetzner. La Compañía reembolsará los costos documentados de terceros atribuibles al uso de la Compañía, sin margen, conforme a las reglas de aprobación y reembolso. El alojamiento no cambia la propiedad. La Compañía recibirá acceso administrativo vigente, credenciales, capacidad de exportación y copias de seguridad razonablemente actuales. SR cooperará con una migración aprobada por Voto Mayoritario y no suspenderá acceso para presionar a la Compañía o a otro Miembro. Tras solicitarse la migración o terminar la participación de SR, prestará asistencia razonable durante un máximo de 30 días."),
        ("7.6 Información Confidencial", "Información Confidencial incluye planes, precios, finanzas, credenciales, código fuente, automatizaciones, prompts, información de clientes y prospectos, estrategias, analítica, secretos comerciales, datos personales y anexos no públicos. Un Miembro solo la usará para fines de la Compañía y solo la revelará a personas autorizadas sujetas a obligaciones adecuadas o cuando la ley lo exija, previa notificación cuando esté permitida."),
        ("7.7 Seguridad y Devolución", "Cada Miembro aplicará medidas razonables de seguridad, informará prontamente una posible vulneración y devolverá o eliminará de forma segura la información y credenciales cuando se solicite o al separarse, salvo una copia archivada exigida por ley."),
        ("7.8 Vigencia Posterior", "La confidencialidad continuará indefinidamente para secretos comerciales y durante cinco años después de la separación para otra Información Confidencial, o por más tiempo si una ley imperativa o contrato con cliente lo exige."),
    ]),
    ("Artículo 8 Emisión Admisión y Transferencias", [
        ("8.1 Nuevas Participaciones y Miembros", "La Compañía solo podrá emitir una participación o admitir un nuevo Miembro con aprobación escrita unánime. El nuevo Miembro firmará la adhesión del Anexo D o documento sustancialmente similar. Un cesionario solo recibirá derechos económicos salvo admisión unánime como Miembro."),
        ("8.2 Derechos de Suscripción Preferente", "Antes de emitir una nueva participación, opción, warrant o derecho convertible, la Compañía ofrecerá a cada Miembro comprar su parte proporcional en las mismas condiciones materiales. La oferta permanecerá abierta 20 días hábiles. Una transacción aprobada por unanimidad podrá establecer procedimientos distintos, pero la renuncia de un Miembro exigirá su consentimiento escrito."),
        ("8.3 Restricción de Transferencia", "Ningún Miembro podrá vender, ceder, donar, pignorar, gravar o transferir una participación salvo conforme a este Acuerdo. Una transferencia no autorizada será ineficaz en la máxima medida legal."),
        ("8.4 Derecho de Preferencia de Compra", "Antes de transferir voluntariamente a un tercero, el Miembro entregará a la Compañía y a los demás Miembros la oferta escrita bona fide completa. La Compañía tendrá 20 días hábiles para comprar todo o parte en esas condiciones. Los demás Miembros tendrán luego 15 días hábiles para comprar el saldo proporcionalmente o según acuerden. Si los derechos no se ejercen totalmente, la transferencia podrá cerrarse dentro de 60 días en condiciones no más favorables al tercero, sujeta al derecho de acompañamiento y a que el comprador no será Miembro sin admisión unánime."),
        ("8.5 Derecho de Acompañamiento", "Si un tercero propone comprar una participación de un Miembro, cada otro Miembro podrá incluir la misma proporción de su propia participación en condiciones económicas idénticas por punto porcentual. El vendedor deberá lograr que el comprador adquiera esas participaciones o no podrá completar la transferencia."),
        ("8.6 Sin Venta Forzada por Mayoría", "No existirá derecho de arrastre. La venta de todas las participaciones o de todos o sustancialmente todos los activos requerirá aprobación unánime."),
        ("8.7 Cumplimiento y Retención", "Toda transferencia quedará sujeta a leyes de valores, fiscales, sanciones, prevención de lavado y demás normas aplicables. La Compañía y el comprador podrán retener montos exigidos por ley, incluida la retención relacionada con la transferencia por un extranjero de una participación en una sociedad."),
    ]),
    ("Artículo 9 Separación Remoción y Compra de Participaciones", [
        ("9.1 Retiro Voluntario", "Un Miembro podrá retirarse de la administración mediante aviso escrito con al menos 90 días. En la fecha efectiva, será titular separado de derechos económicos, dejará de votar y de actuar por la Compañía, pero conservará derechos económicos hasta la compra o transferencia de su participación."),
        ("9.2 Opciones de Compra", "Después del retiro voluntario, muerte, Incapacidad Permanente o remoción por Causa Justificada, la Compañía tendrá la primera opción de comprar toda la participación afectada. Podrá ejercerla dentro de 60 días después de recibir aviso del hecho o de la valoración final, lo que ocurra después. Si no compra todo, los demás Miembros tendrán 30 días adicionales para comprar el saldo proporcionalmente o según acuerden."),
        ("9.3 Muerte y Herederos", "Al fallecer un Miembro, su representante o herederos recibirán únicamente derechos económicos y no serán Miembros. Deberán vender si la Compañía o los demás Miembros ejercen oportunamente la opción. Hasta el cierre, recibirán distribuciones atribuibles, sujetas a retenciones y compensaciones."),
        ("9.4 Incapacidad Permanente", "Incapacidad Permanente significa la imposibilidad, por condición física o mental, de participar materialmente en la administración durante 180 días consecutivos, sustentada por evidencia médica razonable que respete la privacidad. Creará una opción de compra, pero no constituirá conducta indebida."),
        ("9.5 Causa Justificada", "Causa Justificada significa fraude, hurto, apropiación indebida, conducta dolosa materialmente perjudicial, uso consciente indebido de fondos o credenciales, incumplimiento material de confidencialidad o propiedad intelectual, condena por delito grave de deshonestidad u otro incumplimiento material no subsanado dentro de 15 días hábiles después de aviso escrito cuando sea subsanable."),
        ("9.6 Procedimiento de Remoción", "Los otros dos Miembros iniciales podrán remover de la administración a un Miembro por Causa Justificada. El afectado recibirá aviso escrito detallado y al menos cinco días hábiles para responder, salvo suspensión inmediata razonablemente necesaria para proteger fondos, sistemas, evidencia o clientes. Podrá presentar información, pero no votar. La remoción termina administración y voto, pero no confisca la propiedad económica. La Compañía podrá exigir la venta conforme a la Sección 9.2."),
        ("9.7 Valor Justo de Mercado", "El precio será la parte proporcional del valor justo de mercado de la Compañía como negocio en marcha, sin descuentos por minoría o falta de mercado, menos únicamente compensaciones acordadas por escrito o establecidas por sentencia o laudo final. La Compañía y el vendedor intentarán acordar el valor dentro de 20 días hábiles. Si no lo hacen, nombrarán conjuntamente un valuador empresarial independiente. Si no acuerdan el valuador en diez días hábiles, el ICDR lo designará a solicitud de cualquiera. Su determinación será vinculante salvo error manifiesto y el costo se dividirá por mitad."),
        ("9.8 Condiciones de Pago", "Salvo otro acuerdo, el comprador pagará 20 por ciento al cierre y el saldo en 24 cuotas mensuales iguales mediante pagaré. El interés será la tasa preferencial de Estados Unidos publicada por The Wall Street Journal en la fecha de cierre más dos puntos porcentuales, sin superar el máximo legal. Podrá pagarse anticipadamente sin penalidad. Ningún Miembro otorgará garantía personal solo por la compra. Un pago prohibido se aplazará hasta ser lícito y seguirá causando intereses."),
        ("9.9 Cierre", "El cierre ocurrirá dentro de 30 días después de determinar el precio, sujeto a retenciones legales. El vendedor transferirá la participación libre de gravámenes, renunciará a cargos, devolverá bienes y firmará documentos razonables. El comprador entregará el pago inicial y el pagaré."),
        ("9.10 Continuidad", "La separación, muerte, incapacidad, quiebra o transferencia de un Miembro no disolverá la Compañía."),
    ]),
    ("Artículo 10 Deberes y Protección", [
        ("10.1 Normas de Conducta", "Cada Miembro actuará de buena fe, tratará justamente a la Compañía y los demás Miembros, evitará daño intencional, revelará conflictos materiales, protegerá bienes y cumplirá los deberes que no puedan renunciarse conforme a la ley de Wyoming."),
        ("10.2 Confianza en Información", "Un Miembro que actúe de buena fe podrá confiar en registros de la Compañía e información, opiniones e informes de profesionales calificados que razonablemente considere fiables."),
        ("10.3 Indemnización", "En la máxima medida legal, la Compañía indemnizará a un Miembro por reclamaciones o gastos derivados de actividad autorizada realizada de buena fe. No habrá indemnización por fraude, hurto, conducta dolosa, violación consciente de ley, beneficio personal no autorizado o incumplimiento material del Acuerdo."),
        ("10.4 Anticipo de Gastos", "La Compañía podrá anticipar gastos de defensa tras recibir promesa escrita de reembolso si finalmente no procede la indemnización. Requerirá aprobación de todos los Miembros sin conflicto."),
        ("10.5 Seguros", "La Compañía podrá contratar seguros de responsabilidad, ciberseguridad, errores y omisiones, persona clave u otros mediante Voto Mayoritario."),
    ]),
    ("Artículo 11 No Captación y Protecciones Comerciales", [
        ("11.1 Sin Prohibición General de Competencia", "Este Acuerdo no impide ampliamente que un antiguo Miembro obtenga ingresos u opere otro negocio. Las obligaciones específicas de confidencialidad, propiedad intelectual y no captación protegen intereses legítimos."),
        ("11.2 No Captación de Clientes", "Durante la membresía y 24 meses después de la separación, un antiguo Miembro no solicitará conscientemente que un cliente termine o reduzca materialmente su relación cuando el antiguo Miembro tuvo contacto material o recibió Información Confidencial material sobre ese cliente durante los doce meses anteriores."),
        ("11.3 No Captación de Personal", "Durante la membresía y 24 meses después de la separación, un antiguo Miembro no solicitará conscientemente que un empleado o contratista con quien trabajó materialmente durante los doce meses anteriores termine su relación. La publicidad general no dirigida y la contratación de quien responda sin solicitud dirigida no violan esta Sección."),
        ("11.4 Alcance Lícito", "Las Secciones 11.2 y 11.3 se aplicarán únicamente en la máxima medida permitida por la ley aplicable a la persona y conducta. Un tribunal o árbitro podrá limitar duración, actividad o alcance inválidos en lugar de anular más de lo necesario. Nada restringe denuncias protegidas, competencia lícita o uso de habilidades generales."),
        ("11.5 Medidas Cautelares", "El uso indebido real o amenazado de propiedad intelectual, secretos comerciales, credenciales o Información Confidencial puede causar daños no reparables adecuadamente con dinero. La Compañía podrá solicitar medidas temporales o cautelares conforme a la Sección 13.4 sin renunciar al arbitraje."),
    ]),
    ("Artículo 12 Registros y Cumplimiento", [
        ("12.1 Registros", "La Compañía conservará documentos constitutivos, este Acuerdo y modificaciones, consentimientos, estados financieros, declaraciones fiscales, contratos materiales, registros de propiedad y demás registros exigidos. Se permiten registros electrónicos."),
        ("12.2 Inspección", "Un Miembro podrá inspeccionar registros razonables para un fin relacionado con su membresía, con aviso razonable y sujeto a confidencialidad, privilegio, privacidad y seguridad. Un titular separado recibirá solo la información exigida por ley o necesaria para verificar distribuciones y una compra."),
        ("12.3 Cuentas Separadas y Documentación", "Los Miembros preservarán la existencia jurídica separada, utilizarán cuentas y contratos de la Compañía, documentarán aprobaciones materiales y evitarán representar que obligaciones de la Compañía son personales salvo garantía separada firmada intencionalmente."),
        ("12.4 Cumplimiento Multijurisdiccional", "Antes de abrir oficina, contratar personal, crear establecimiento habitual o dar otro paso que probablemente exija registro o presencia fiscal en un país o estado, los Miembros obtendrán asesoría y aprobarán la medida bajo la regla de voto aplicable. Cada Miembro informará correctamente dónde presta servicios."),
    ]),
    ("Artículo 13 Resolución de Controversias", [
        ("13.1 Conversación de Buena Fe", "Una parte primero dará aviso escrito describiendo la controversia y la solución solicitada. Las partes se reunirán remotamente dentro de diez días hábiles e intentarán resolverla de buena fe."),
        ("13.2 Mediación", "Si no se resuelve, participarán en mediación remota confidencial con mediador elegido conjuntamente. Si no lo eligen en diez días hábiles, el ICDR lo proporcionará o designará. Compartirán los honorarios por mitad salvo acuerdo distinto. La mediación terminará 30 días después del nombramiento salvo prórroga escrita."),
        ("13.3 Arbitraje", "Toda controversia no resuelta se decidirá definitivamente por un árbitro conforme a los Procedimientos Internacionales de Resolución de Controversias del International Centre for Dispute Resolution y sus reglas comerciales aplicables. La sede jurídica será Cheyenne, Wyoming. Las audiencias serán remotas salvo decisión del árbitro. El idioma será inglés. El laudo podrá otorgar remedios legales o equitativos y asignar honorarios y costos legales razonables. Podrá ejecutarse ante cualquier tribunal competente."),
        ("13.4 Medidas Judiciales Urgentes", "Una parte podrá solicitar medidas temporales a un tribunal competente para proteger fondos, sistemas, credenciales, evidencia, Información Confidencial, propiedad intelectual o datos de clientes mientras se tramita el arbitraje. Esto no renuncia a la mediación o arbitraje."),
        ("13.5 Confidencialidad", "Las partes mantendrán confidenciales la mediación y el arbitraje salvo cuando sea necesario para ejecutar el laudo, obtener asesoría profesional, cumplir la ley o proteger un derecho."),
    ]),
    ("Artículo 14 Disolución y Liquidación", [
        ("14.1 Causales de Disolución", "La Compañía se disolverá por aprobación escrita unánime, por un hecho que vuelva ilícito el negocio o por sentencia final que ordene disolución. El bloqueo decisorio por sí solo no será causal."),
        ("14.2 Liquidación", "Una persona elegida por unanimidad liquidará la Compañía. Recaudará activos, completará o terminará contratos, pagará o reservará pasivos y distribuirá el remanente conforme a la ley y este Acuerdo."),
        ("14.3 Distribución Final", "Después de satisfacer acreedores y reservas, los activos restantes se distribuirán conforme a cuentas de capital positivas cuando lo exija la ley fiscal y, en otro caso, según las Participaciones Porcentuales, sujeto a compensaciones y retenciones legales."),
    ]),
    ("Artículo 15 Disposiciones Generales", [
        ("15.1 Avisos", "Los avisos formales se enviarán a las direcciones postales y electrónicas del Anexo A, actualizadas mediante aviso escrito. El correo electrónico será efectivo al enviarse sin aviso de fallo de entrega. Un aviso de Emergencia se enviará además por el canal grupal designado. La notificación judicial seguirá la ley aplicable."),
        ("15.2 Modificaciones", "Este Acuerdo solo podrá modificarse mediante escrito firmado por todos los Miembros. Ninguna conducta o conversación electrónica lo modificará salvo que declare claramente intención de modificar y reciba todas las firmas exigidas."),
        ("15.3 Acuerdo Completo", "Este Acuerdo, sus anexos y toda adhesión o modificación firmada constituyen el acuerdo completo sobre gobierno y propiedad. Sustituyen entendimientos anteriores incompatibles."),
        ("15.4 Prelación", "Entre los Miembros, este Acuerdo prevalecerá en la máxima medida legal. Un documento presentado ante el Estado prevalecerá solo cuando la ley le otorgue prioridad. Una norma imperativa prevalecerá sobre una disposición incompatible."),
        ("15.5 Divisibilidad", "Si una disposición no es exigible, se aplicará en la máxima medida lícita o se separará, y el resto continuará. Un tribunal o árbitro podrá reformar una restricción solo cuando la ley lo permita."),
        ("15.6 Renuncia", "Toda renuncia será escrita y se aplicará solo al caso indicado. La demora o ejercicio parcial de un derecho no constituye renuncia."),
        ("15.7 Sin Beneficiarios Terceros", "Salvo una persona indemnizada y un sucesor permitido, este Acuerdo no crea derechos en terceros."),
        ("15.8 Ejemplares y Firmas Electrónicas", "Este Acuerdo podrá firmarse en ejemplares separados y mediante un método fiable de firma electrónica. Todos formarán un solo instrumento."),
        ("15.9 Documentos Adicionales", "Cada Miembro firmará documentos razonables para ejecutar este Acuerdo, incluidos documentos fiscales, bancarios, de transferencia y propiedad intelectual."),
        ("15.10 Prevalencia del Inglés", "La versión en español se proporciona para conveniencia y comprensión mutua. Si existe diferencia, prevalecerá la versión en inglés."),
    ]),
]


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=100, start=100, bottom=100, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr, fld_char2])


def add_clause(doc, label, text):
    p = doc.add_paragraph(style="Body Text")
    p.paragraph_format.keep_together = True
    p.paragraph_format.widow_control = True
    r = p.add_run(label + "  ")
    r.bold = True
    p.add_run(text)


def add_part(doc, title, intro, articles):
    h = doc.add_paragraph(title, style="Title")
    h.paragraph_format.page_break_before = True
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph(intro, style="Body Text")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    for heading, clauses in articles:
        doc.add_heading(heading, level=1)
        for label, text in clauses:
            add_clause(doc, label, text)


def add_signature_line(doc, name, country):
    block = []
    p = doc.add_paragraph(style="Body Text")
    block.append(p)
    p.paragraph_format.space_before = Pt(18)
    p.add_run("Signature Firma  ").bold = True
    p.add_run("____________________________________________")
    p = doc.add_paragraph(style="Body Text")
    block.append(p)
    p.add_run("Name Nombre  ").bold = True
    p.add_run(name)
    p = doc.add_paragraph(style="Body Text")
    block.append(p)
    p.add_run("Country of residence País de residencia  ").bold = True
    p.add_run(country)
    p = doc.add_paragraph(style="Body Text")
    block.append(p)
    p.add_run("Date Fecha  ").bold = True
    p.add_run("____________________________________________")
    for paragraph in block[:-1]:
        paragraph.paragraph_format.keep_with_next = True


def add_bilingual_exhibit_heading(doc, english, spanish):
    h = doc.add_paragraph(english, style="Title")
    h.paragraph_format.page_break_before = True
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph(spanish, style="Subtitle")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_bilingual_clause(doc, label, en, es):
    add_clause(doc, label, en)
    p = doc.add_paragraph(es, style="Spanish Body")
    p.paragraph_format.left_indent = Inches(0.2)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.9)
section.left_margin = Inches(0.85)
section.right_margin = Inches(0.85)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Liberation Serif"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Liberation Serif")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Liberation Serif")
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.04

for style_name in ("Body Text", "Spanish Body"):
    if style_name not in styles:
        styles.add_style(style_name, 1)
    style = styles[style_name]
    style.base_style = styles["Normal"]
    style.font.name = "Liberation Serif"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Liberation Serif")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Liberation Serif")
    style.font.size = Pt(10.5)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_after = Pt(5)
    style.paragraph_format.line_spacing = 1.04

title_style = styles["Title"]
title_style.font.name = "Liberation Sans"
title_style._element.rPr.rFonts.set(qn("w:ascii"), "Liberation Sans")
title_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Liberation Sans")
title_style.font.size = Pt(20.5)
title_style.font.bold = True
title_style.font.color.rgb = RGBColor(0, 0, 0)
title_style.paragraph_format.space_after = Pt(12)
title_style.paragraph_format.space_before = Pt(10)
title_ppr = title_style._element.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)

subtitle_style = styles["Subtitle"]
subtitle_style.font.name = "Liberation Sans"
subtitle_style._element.rPr.rFonts.set(qn("w:ascii"), "Liberation Sans")
subtitle_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Liberation Sans")
subtitle_style.font.size = Pt(13)
subtitle_style.font.color.rgb = RGBColor(0, 0, 0)

for level, size in ((1, 14), (2, 11.5)):
    style = styles[f"Heading {level}"]
    style.font.name = "Liberation Sans"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Liberation Sans")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Liberation Sans")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.keep_with_next = True
    style.paragraph_format.space_before = Pt(13 if level == 1 else 8)
    style.paragraph_format.space_after = Pt(6)

header = section.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.CENTER
hr = header.add_run("RANKING REBELS LLC   OPERATING AGREEMENT   ACUERDO OPERATIVO")
hr.font.name = "Liberation Sans"
hr.font.size = Pt(8)
hr.font.color.rgb = RGBColor(80, 80, 80)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = footer.add_run("Ranking Rebels LLC   ")
fr.font.name = "Liberation Sans"
fr.font.size = Pt(8)
add_page_number(footer)

# Cover
p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(100)
p.add_run("Operating Agreement of Ranking Rebels LLC")
p = doc.add_paragraph(style="Subtitle")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("Acuerdo Operativo de Ranking Rebels LLC")
p = doc.add_paragraph(style="Body Text")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(24)
p.add_run("Wyoming member managed limited liability company\nCompañía de responsabilidad limitada de Wyoming administrada por sus Miembros")
p = doc.add_paragraph(style="Body Text")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(36)
r = p.add_run("Effective date Fecha de vigencia\n[FORMATION DATE  FECHA DE CONSTITUCIÓN]")
r.bold = True

checklist_heading = doc.add_heading("Completion Before Signature", level=1)
checklist_heading.paragraph_format.page_break_before = True
for item in [
    "Insert the formation date, principal office, registered agent, Member addresses and emails, and the eligible United States partnership representative.",
    "Confirm that the filed Articles of Organization identify the Company as member-managed or do not identify it as manager-managed.",
    "Have a Wyoming attorney and advisers familiar with United States, Australian, Dutch, and Colombian tax review the completed document.",
    "Each Member should sign the Agreement and Initial Unanimous Consent. SR and an authorized Company representative should also sign the Domain and Intellectual Property Assignment.",
    "Keep the signed Agreement in the Company's private records. It is not filed with the Wyoming Secretary of State.",
]:
    p = doc.add_paragraph(style="Body Text")
    p.style = styles["Body Text"]
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.first_line_indent = Inches(-0.18)
    p.add_run("•  ").bold = True
    p.add_run(item)

doc.add_heading("Antes de Firmar", level=1)
for item in [
    "Complete la fecha de constitución, oficina principal, agente registrado, direcciones y correos de los Miembros, y el representante fiscal elegible en Estados Unidos.",
    "Confirme que los Artículos de Organización indican que la Compañía es administrada por sus Miembros o no indican que sea administrada por gerentes.",
    "Solicite revisión de un abogado de Wyoming y asesores fiscales con experiencia en Estados Unidos, Australia, Países Bajos y Colombia.",
    "Cada Miembro debe firmar el Acuerdo y el Consentimiento Unánime Inicial. SR y un representante autorizado también deben firmar la Cesión de Dominio y Propiedad Intelectual.",
    "Conserve el Acuerdo firmado en los registros privados de la Compañía. No se presenta ante la Secretaría de Estado de Wyoming.",
]:
    p = doc.add_paragraph(style="Spanish Body")
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.first_line_indent = Inches(-0.18)
    p.add_run("•  ").bold = True
    p.add_run(item)

add_part(doc, "Part One English Version", "This version contains the controlling terms of the Agreement.", ENGLISH)
add_part(doc, "Parte Dos Versión en Español", "Esta versión facilita la comprensión. La versión en inglés prevalece si existe alguna diferencia.", SPANISH)

# Signature page
add_bilingual_exhibit_heading(doc, "Signature Page", "Página de Firmas")
p = doc.add_paragraph(style="Body Text")
p.add_run("By signing below, each Member adopts both language versions of this Agreement and agrees that the English version controls.\n").bold = True
p.add_run("Al firmar, cada Miembro adopta ambas versiones lingüísticas de este Acuerdo y acepta que la versión en inglés prevalece.")
for name, country in [
    ("Juan David Arteaga Henriques", "Australia"),
    ("Santiago Mejia Montoya", "Colombia"),
    ("Santiago Ramirez Castaño", "Netherlands Países Bajos"),
]:
    add_signature_line(doc, name, country)

# Schedule A
add_bilingual_exhibit_heading(doc, "Schedule A Members and Contributions", "Anexo A Miembros y Aportes")
p = doc.add_paragraph("Member contact information appears only in this Schedule. Update it through written notice under Section 15.1.\nLos datos de contacto de los Miembros aparecen únicamente en este Anexo. Se actualizan mediante aviso escrito conforme a la Sección 15.1.", style="Body Text")
table = doc.add_table(rows=1, cols=5)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.style = "Table Grid"
widths = [1.42, 0.85, 0.85, 1.72, 2.0]
headers = ["Member Miembro", "Residence Residencia", "Interest Participación", "Address and Email Dirección y Correo", "Acknowledged Contribution Aporte Reconocido"]
for i, (cell, text) in enumerate(zip(table.rows[0].cells, headers)):
    cell.text = text
    set_cell_shading(cell, "404040")
    set_cell_margins(cell, top=55, start=65, bottom=55, end=65)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    cell.width = Inches(widths[i])
    for run in cell.paragraphs[0].runs:
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.bold = True
        run.font.size = Pt(7.2)
set_repeat_table_header(table.rows[0])
rows = [
    ("Juan David Arteaga Henriques  JD", "Australia", "33 1/3%", "[MAILING ADDRESS]\n[EMAIL]", "Marketing relationships, access to Australian markets, and introduction to one prospective customer. No signed customer contract exists at formation."),
    ("Santiago Mejia Montoya  SM", "Colombia", "33 1/3%", "[MAILING ADDRESS]\n[EMAIL]", "SEO and marketing experience and social proof."),
    ("Santiago Ramirez Castaño  SR", "Netherlands Países Bajos", "33 1/3%", "[MAILING ADDRESS]\n[EMAIL]", "AI and technology experience, access to European markets, assignment of rankingrebels.com, and initial Hetzner hosting subject to reimbursement."),
]
for ridx, row_data in enumerate(rows):
    cells = table.add_row().cells
    for i, (cell, text) in enumerate(zip(cells, row_data)):
        cell.text = text
        set_cell_margins(cell, top=45, start=60, bottom=45, end=60)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell.width = Inches(widths[i])
        if ridx % 2 == 1:
            set_cell_shading(cell, "F2F4F7")
        for p2 in cell.paragraphs:
            for run in p2.runs:
                run.font.size = Pt(6.9)

# Exhibit B
add_bilingual_exhibit_heading(doc, "Exhibit B Domain and Intellectual Property Assignment", "Anexo B Cesión de Dominio y Propiedad Intelectual")
add_bilingual_clause(doc, "B.1 Assignment", "For good and valuable consideration, Santiago Ramirez Castaño assigns to Ranking Rebels LLC all worldwide right, title, and interest in the domain name rankingrebels.com, together with its associated goodwill, DNS configuration under his control, and transferable website materials created specifically for that domain before the effective date.", "Por contraprestación válida y suficiente, Santiago Ramirez Castaño cede a Ranking Rebels LLC todos los derechos mundiales sobre el nombre de dominio rankingrebels.com, junto con su goodwill asociado, la configuración DNS bajo su control y los materiales transferibles del sitio creados específicamente para ese dominio antes de la fecha efectiva.")
add_bilingual_clause(doc, "B.2 Effective Time", "This assignment becomes effective when Ranking Rebels LLC is formed. Until the registrar transfer is complete, SR will maintain the registration solely for the Company's benefit and will not sell, pledge, redirect, or disable the domain.", "Esta cesión entrará en vigor cuando se constituya Ranking Rebels LLC. Hasta completar la transferencia en el registrador, SR mantendrá el registro únicamente en beneficio de la Compañía y no venderá, gravará, redirigirá ni deshabilitará el dominio.")
add_bilingual_clause(doc, "B.3 Further Action", "SR will provide authorization codes, credentials, signatures, and reasonable cooperation needed to transfer and secure the domain in an account controlled by the Company. The Company will pay reasonable registrar transfer fees.", "SR entregará códigos de autorización, credenciales, firmas y cooperación razonable para transferir y proteger el dominio en una cuenta controlada por la Compañía. La Compañía pagará los costos razonables del registrador.")
add_bilingual_clause(doc, "B.4 Retained Rights", "SR retains general knowledge, experience, and material not created specifically for Ranking Rebels and not incorporated into Company property. No ownership right in the assigned domain or associated goodwill is retained.", "SR conserva conocimientos generales, experiencia y materiales no creados específicamente para Ranking Rebels ni incorporados a bienes de la Compañía. No conserva derecho de propiedad sobre el dominio cedido o su goodwill asociado.")
add_signature_line(doc, "Santiago Ramirez Castaño  Assignor Cedente", "Netherlands Países Bajos")
add_signature_line(doc, "Ranking Rebels LLC  By authorized Member Por Miembro autorizado", "Wyoming United States")

# Exhibit C
add_bilingual_exhibit_heading(doc, "Exhibit C Initial Unanimous Written Consent", "Anexo C Consentimiento Escrito Unánime Inicial")
p = doc.add_paragraph(style="Body Text")
p.add_run("Effective on the Company's formation date, the undersigned unanimously approve the following actions.\n").bold = True
p.add_run("Con efecto en la fecha de constitución, los firmantes aprueban unánimemente las siguientes acciones.")
consents = [
    ("Adoption", "The bilingual Operating Agreement is adopted, with the English version controlling.", "Se adopta el Acuerdo Operativo bilingüe, prevaleciendo la versión en inglés."),
    ("Membership", "The three persons in Schedule A are admitted as Members with 33 1/3 percent each, without vesting.", "Las tres personas del Anexo A son admitidas como Miembros con 33 1/3 por ciento cada una, sin consolidación gradual."),
    ("Assets", "The Domain and Intellectual Property Assignment is accepted, and all future customer contracts will be entered in the Company's name.", "Se acepta la Cesión de Dominio y Propiedad Intelectual, y los contratos futuros con clientes se celebrarán a nombre de la Compañía."),
    ("Formation Documents", "The filed Articles of Organization and registered-agent appointment are ratified.", "Se ratifican los Artículos de Organización presentados y el nombramiento del agente registrado."),
    ("EIN and Banking", "Any Member may complete an EIN application and banking onboarding using accurate information, but account controls and expenditures remain subject to the Agreement.", "Cualquier Miembro podrá completar la solicitud de EIN y apertura bancaria con información exacta, sujetos los controles y gastos al Acuerdo."),
    ("Tax Professional", "The Company is authorized to engage a United States international-tax professional and appoint an eligible United States partnership representative after unanimous approval.", "La Compañía queda autorizada para contratar un profesional fiscal internacional en Estados Unidos y nombrar un representante fiscal elegible tras aprobación unánime."),
    ("Preformation Expenses", "Documented preformation expenses will be reimbursed only after Majority Vote and subject to available Company funds.", "Los gastos documentados anteriores a la constitución solo se reembolsarán mediante Voto Mayoritario y sujetos a fondos disponibles."),
]
for label, en, es in consents:
    add_bilingual_clause(doc, label, en, es)
doc.add_page_break()
for name, country in [
    ("Juan David Arteaga Henriques", "Australia"),
    ("Santiago Mejia Montoya", "Colombia"),
    ("Santiago Ramirez Castaño", "Netherlands Países Bajos"),
]:
    add_signature_line(doc, name, country)

# Exhibit D
add_bilingual_exhibit_heading(doc, "Exhibit D New Member Joinder", "Anexo D Adhesión de Nuevo Miembro")
add_bilingual_clause(doc, "D.1 Joinder", "The undersigned joins the Operating Agreement of Ranking Rebels LLC as a Member, agrees to be bound by it as amended, and confirms receipt of both language versions. The undersigned acknowledges that the English version controls.", "El firmante se adhiere al Acuerdo Operativo de Ranking Rebels LLC como Miembro, acepta quedar obligado por sus modificaciones y confirma haber recibido ambas versiones. Reconoce que prevalece la versión en inglés.")
add_bilingual_clause(doc, "D.2 Admission", "This joinder is effective only when every existing Member gives the unanimous written approval required by the Agreement and the Company updates Schedule A.", "Esta adhesión solo será efectiva cuando todos los Miembros existentes otorguen la aprobación escrita unánime exigida y la Compañía actualice el Anexo A.")
for label in ["New Member Nuevo Miembro", "Ranking Rebels LLC By Por", "Existing Members Approval Aprobación de Miembros Existentes"]:
    p = doc.add_paragraph(style="Body Text")
    p.paragraph_format.space_before = Pt(10)
    p.add_run(label + "\n").bold = True
    p.add_run("Name Nombre  ____________________________________________\n")
    p.add_run("Signature Firma  _________________________________________\n")
    p.add_run("Date Fecha  _____________________________________________")

# Core properties and final save
doc.core_properties.title = "Operating Agreement of Ranking Rebels LLC"
doc.core_properties.subject = "Bilingual English and Spanish Wyoming LLC operating agreement"
doc.core_properties.author = "Ranking Rebels LLC"
doc.core_properties.keywords = "Wyoming LLC operating agreement bilingual English Spanish"

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT.resolve())
